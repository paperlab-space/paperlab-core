from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from paperlab_core.config import RunConfig
from paperlab_core.export import latex_table, write_all
from paperlab_core.export.latex import escape
from paperlab_core.results import RunResults
from paperlab_core.runner import run

from .fakes import FakeBackend, ProviderError


@pytest.fixture
def results(multilingual_xlsx: Path) -> RunResults:
    config = RunConfig.from_dict(
        {
            "dataset": {"path": str(multilingual_xlsx), "sheet": "test"},
            "mapping": {"input": "source", "expected": "gold", "id": "id"},
            "prompt": "Translate: {source}",
            "models": [{"id": "openai/good_model", "label": "Good_Model"}, "openai/bad"],
            "metrics": ["exact_match", "f1"],
            "seed": 1,
        }
    )

    def answer(model: object, prompt: str, attempt: int) -> str | Exception:
        if getattr(model, "id", "") == "openai/bad":
            return ProviderError("boom") if "你好" in prompt else '=HYPERLINK("x")'
        return "thank you" if any(s in prompt for s in ("ありがとう", "감사")) else "hello"

    return run(config, backend=FakeBackend(answer))


def test_write_all(results: RunResults, tmp_path: Path) -> None:
    paths = write_all(results, tmp_path / "out")
    assert {p.name for p in paths.values()} == {
        "results.xlsx",
        "table.tex",
        "run.yaml",
        "results.json",
    }

    workbook = load_workbook(paths["excel"])
    assert workbook.sheetnames == ["Summary", "Outputs", "Config"]
    summary = list(workbook["Summary"].iter_rows(values_only=True))
    assert summary[0][:4] == ("Model", "Model ID", "exact_match", "f1")
    assert summary[1][0] == "Good_Model" and summary[1][2] == 1.0
    assert summary[2][0] == "bad" and summary[2][2] == 0.0 and summary[2][5] == 1  # failed
    outputs = list(workbook["Outputs"].iter_rows(values_only=True))
    assert outputs[0] == (
        "id",
        "source",
        "expected",
        "Good_Model | output",
        "Good_Model | exact_match",
        "bad | output",
        "bad | exact_match",
        "bad | error",
    )
    assert outputs[1][:5] == ("1", "မင်္ဂလာပါ", "hello", "hello", 1)
    assert outputs[3][7] == "boom"
    formula_cell = workbook["Outputs"]["F2"]
    assert formula_cell.value == '=HYPERLINK("x")' and formula_cell.data_type == "s"

    repro = RunConfig.from_yaml(paths["repro"])  # loads back as a config
    assert repro.dataset is not None and Path(repro.dataset.path).name == "data.xlsx"
    text = paths["repro"].read_text(encoding="utf-8")
    assert "sha256" in text and str(tmp_path) not in text  # no local paths leak

    again = RunResults.load_json(paths["json"])
    assert again.scores["Good_Model"]["f1"].value == results.scores["Good_Model"]["f1"].value


def test_latex_table(results: RunResults) -> None:
    tex = latex_table(results)
    assert r"\begin{tabular}{lrr}" in tex
    assert r"Model & EM & F1 \\" in tex
    assert r"Good\_Model & \textbf{100.0} & \textbf{100.0} \\" in tex
    assert r"bad & 0.0 & 0.0 \\" in tex
    assert r"\caption{Results on 5 examples (\%).}" in tex
    raw = latex_table(
        results, metrics=["f1"], percent=False, decimals=3, floating=False, bold_best=False
    )
    assert "Good\\_Model & 1.000 \\\\" in raw and r"\begin{table}" not in raw
    with pytest.raises(ValueError, match="Unknown metric"):
        latex_table(results, metrics=["bleu"])


def test_escape() -> None:
    assert escape("a_b & 50% #1 {x} ~^\\") == (
        r"a\_b \& 50\% \#1 \{x\} \textasciitilde{}\textasciicircum{}\textbackslash{}"
    )
