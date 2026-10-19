from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pytest

from paperlab_space.config import MetricSpec, ModelSpec, RunConfig
from paperlab_space.data import Dataset
from paperlab_space.errors import ConfigError
from paperlab_space.export import latex_table, write_all
from paperlab_space.metrics import get_metric
from paperlab_space.results import RunResults
from paperlab_space.runner import run
from paperlab_space.stats import adjust, compare, confidence_interval

from .fakes import FakeBackend

N = 60
DATA = Dataset.from_records(
    [{"id": str(i), "q": f"q{i}", "a": "yes" if i % 2 else "no"} for i in range(N)]
)


def respond(model: ModelSpec, prompt: str, attempt: int) -> str:
    i = int(prompt[1:])
    gold = "yes" if i % 2 else "no"
    wrong = "no" if gold == "yes" else "yes"
    if model.id == "openai/strong":
        return gold if i % 10 else wrong  # 90% correct
    if model.id == "openai/twin":
        return gold if (i + 1) % 10 else wrong  # also 90%, different rows
    return gold if i % 2 == 0 else wrong  # 50%


def make_config(**statistics: object) -> RunConfig:
    return RunConfig.from_dict(
        {
            "mapping": {"input": "q", "expected": "a", "id": "id"},
            "prompt": "{q}",
            "models": ["openai/strong", "openai/twin", "openai/weak"],
            "metrics": ["accuracy", "f1"],
            "seed": 7,
            "statistics": statistics,
        }
    )


def test_intervals_on_every_score() -> None:
    results = run(make_config(), DATA, backend=FakeBackend(respond))
    for name in ("strong", "weak"):
        for key in ("accuracy", "f1"):
            result = results.scores[name][key]
            assert result.ci_low is not None and result.ci_high is not None
            assert result.value is not None
            assert result.ci_low <= result.value <= result.ci_high
    strong = results.scores["strong"]["accuracy"]
    assert strong.ci_high is not None and strong.ci_low is not None
    assert 0.05 < strong.ci_high - strong.ci_low < 0.3
    # Deterministic for the same seed.
    again = run(make_config(), DATA, backend=FakeBackend(respond))
    assert again.scores["strong"]["f1"].ci_low == results.scores["strong"]["f1"].ci_low


def test_permutation_test() -> None:
    results = run(make_config(), DATA, backend=FakeBackend(respond))
    by_pair = {(c.metric, c.model_a, c.model_b): c for c in results.comparisons}
    assert len(by_pair) == 6  # 3 pairs x 2 metrics
    clear = by_pair[("accuracy", "strong", "weak")]
    assert clear.difference == pytest.approx(0.4)
    assert clear.p_value is not None and clear.p_value < 0.01 and clear.significant
    assert clear.method == "permutation" and clear.ci_low is None
    same = by_pair[("accuracy", "strong", "twin")]
    assert same.difference == pytest.approx(0.0)
    assert same.p_value is not None and same.p_value > 0.5 and not same.significant
    assert clear.p_adjusted is not None and clear.p_adjusted >= clear.p_value  # Holm


def test_bootstrap_test_and_baseline() -> None:
    config = make_config(significance="bootstrap", baseline="weak", correction="none")
    results = run(config, DATA, backend=FakeBackend(respond))
    assert {(c.model_a, c.model_b) for c in results.comparisons} == {
        ("strong", "weak"),
        ("twin", "weak"),
    }
    c = next(c for c in results.comparisons if c.metric == "accuracy" and c.model_a == "strong")
    assert c.significant and c.p_adjusted == c.p_value
    assert c.ci_low is not None and c.ci_high is not None and 0 < c.ci_low < 0.4 < c.ci_high


def test_statistics_off() -> None:
    results = run(
        make_config(intervals=False, significance="none"), DATA, backend=FakeBackend(respond)
    )
    assert results.comparisons == []
    assert results.scores["strong"]["accuracy"].ci_low is None


def test_unknown_baseline() -> None:
    with pytest.raises(ConfigError, match="statistics.baseline"):
        make_config(baseline="nope")


def test_corpus_metric_comparison() -> None:
    texts = [f"the answer to question {i} is {i * 7} exactly" for i in range(30)]
    data = Dataset.from_records([{"id": str(i), "q": f"q{i}", "a": t} for i, t in enumerate(texts)])

    def answer(model: ModelSpec, prompt: str, attempt: int) -> str:
        i = int(prompt[1:])
        return texts[i] if model.id == "openai/copy" else f"maybe {i}"

    config = RunConfig.from_dict(
        {
            "mapping": {"input": "q", "expected": "a", "id": "id"},
            "prompt": "{q}",
            "models": ["openai/copy", "openai/guess"],
            "metrics": ["bleu", "chrf", "rougeL"],
            "statistics": {"resamples": 200},
        }
    )
    results = run(config, data, backend=FakeBackend(answer))
    bleu = results.scores["copy"]["bleu"]
    assert bleu.value == pytest.approx(1.0) and bleu.ci_low == pytest.approx(1.0)
    assert all(c.significant for c in results.comparisons)


def test_holm_and_bonferroni() -> None:
    p = [0.01, 0.04, 0.03, None]
    assert adjust(p, "bonferroni") == pytest.approx([0.03, 0.12, 0.09, None])
    # Holm: sorted 0.01*3, 0.03*2, 0.04*1, kept monotone.
    assert adjust(p, "holm") == pytest.approx([0.03, 0.06, 0.06, None])
    assert adjust(p, "none") == p


def test_f1_scorer_matches_compute() -> None:
    predictions: list[str | None] = ["a", "b", "c", "a", None, "zzz", "b"]
    references: list[list[str] | None] = [["a"], ["a"], ["c"], ["b"], ["a"], ["c"], None]
    rng = np.random.default_rng(0)
    for params in ({}, {"average": "micro"}, {"average": "weighted"}, {"labels": ["a", "b", "c"]}):
        metric = get_metric(MetricSpec(name="f1", params=params))
        scorer = metric.scorer(predictions, references)
        for _ in range(20):
            sample = rng.integers(0, len(predictions), len(predictions))
            expected = metric.compute(
                [predictions[i] for i in sample], [references[i] for i in sample]
            ).value
            assert scorer(sample) == pytest.approx(expected)


def test_confidence_interval_edge_cases() -> None:
    assert confidence_interval(lambda idx: None, 5) is None
    assert confidence_interval(lambda idx: 1.0, 0) is None
    assert confidence_interval(lambda idx: 0.5, 5) == (0.5, 0.5)


def test_saved_and_exported(tmp_path: Path) -> None:
    config = make_config(baseline="weak")
    results = run(config, DATA, backend=FakeBackend(respond))
    paths = write_all(results, tmp_path)

    loaded = RunResults.load_json(paths["json"])
    assert loaded.comparisons == results.comparisons
    assert loaded.scores["strong"]["f1"].ci_low == results.scores["strong"]["f1"].ci_low
    assert compare(loaded) == results.comparisons

    from openpyxl import load_workbook

    workbook = load_workbook(paths["excel"])
    assert workbook.sheetnames == ["Summary", "Outputs", "Significance", "Config"]
    header = next(workbook["Summary"].iter_rows(values_only=True))
    assert header[2:5] == ("accuracy", "accuracy 95% CI low", "accuracy 95% CI high")
    rows = list(workbook["Significance"].iter_rows(values_only=True))
    assert rows[0][:3] == ("Metric", "Model A", "Model B") and len(rows) == 5

    tex = latex_table(results, ci=True)
    assert r"\textbf{90.0}$^\dagger$ {\scriptsize[" in tex
    assert r"significantly different from weak (paired permutation test, Holm-corrected" in tex
    assert "Brackets: 95\\% bootstrap confidence intervals." in tex
    assert "dagger" not in latex_table(results, mark_significant=False)


def test_old_results_load(tmp_path: Path) -> None:
    results = run(make_config(), DATA, backend=FakeBackend(respond))
    data = results.to_dict()
    del data["comparisons"], data["judgments"]
    for metrics in data["scores"].values():
        for result in metrics.values():
            del result["ci_low"], result["ci_high"], result["info"]
    del data["config"]["statistics"]
    path = tmp_path / "old.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    loaded = RunResults.load_json(path)
    assert loaded.comparisons == [] and loaded.judgments == {}
    assert loaded.scores["strong"]["accuracy"].ci_low is None
