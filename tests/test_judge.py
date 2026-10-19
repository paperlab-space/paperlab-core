from __future__ import annotations

from pathlib import Path

import pytest

from paperlab_space.config import MetricSpec, ModelSpec, RunConfig
from paperlab_space.data import Dataset
from paperlab_space.errors import ConfigError
from paperlab_space.export import write_all
from paperlab_space.metrics import LLMJudge, get_metric, score
from paperlab_space.results import RunResults
from paperlab_space.runner import Progress, estimate, run

from .fakes import FakeBackend, ProviderError

DATA = Dataset.from_records(
    [
        {"id": "r1", "q": "capital of France?", "a": "Paris"},
        {"id": "r2", "q": "capital of Japan?", "a": "Tokyo"},
        {"id": "r3", "q": "capital of Peru?", "a": "Lima"},
        {"id": "r4", "q": "capital of Chad?", "a": ""},
    ]
)
JUDGE = "openai/judge"


def make_config(judge: dict[str, object] | None = None, **overrides: object) -> RunConfig:
    base: dict[str, object] = {
        "mapping": {"input": "q", "expected": "a", "id": "id"},
        "prompt": "Answer: {q}",
        "models": ["openai/good", "openai/bad"],
        "metrics": [{"name": "llm_judge", "model": JUDGE, **(judge or {})}],
        "seed": 3,
    }
    return RunConfig.from_dict({**base, **overrides})


ANSWERS = {"France": "Paris", "Japan": "Tokyo", "Peru": "Lima", "Chad": "N'Djamena"}


def respond(model: ModelSpec, prompt: str, attempt: int) -> str | Exception:
    if model.id == JUDGE:
        # Score 5 when the response matches the reference, else 1.
        response = prompt.split("[Response]\n")[1].split("\n")[0]
        if "[Reference answer]" not in prompt:
            return "Looks fine.\nScore: 4"
        reference = prompt.split("[Reference answer]\n")[1].split("\n")[0]
        return f"Reasoning.\nScore: {5 if response == reference else 1}"
    country = prompt.split("of ")[1].rstrip("?")
    if model.id == "openai/bad":
        return ProviderError("down") if country in ("Peru", "Chad") else "London"
    return ANSWERS[country]


def test_judge_with_reference() -> None:
    backend = FakeBackend(respond)
    judge_progress: list[Progress] = []
    results = run(make_config(), DATA, backend=backend, on_judge_progress=judge_progress.append)

    good = results.scores["good"]["llm_judge"]
    bad = results.scores["bad"]["llm_judge"]
    assert good.per_row == [1.0, 1.0, 1.0, None]  # r4 has no reference
    assert good.value == 1.0 and good.extra["mean_rating"] == 5.0
    # The failed row counts as the lowest score, but not the failed row without a reference.
    assert bad.per_row == [0.0, 0.0, 0.0, None]
    assert bad.extra["failed"] == 1 and bad.info["judge"] == JUDGE

    judge_calls = [prompt for model_id, prompt in backend.calls if model_id == JUDGE]
    assert len(judge_calls) == 5  # 3 rows for good, 2 for bad (one failed)
    assert all("[Task]\nAnswer: capital of" in p for p in judge_calls)
    assert {p.metric for p in judge_progress} == {"llm_judge"}
    assert judge_progress[-1].done == judge_progress[-1].total == 5

    rows = results.judgments["llm_judge"]["good"]
    assert [r.output is not None for r in rows] == [True, True, True, False]


def test_rubric_only_and_parsing() -> None:
    config = make_config({"rubric": "Is it polite? {min_score}-{max_score}", "reference": False})
    results = run(config, DATA, backend=FakeBackend(respond))
    assert results.scores["good"]["llm_judge"].per_row == [0.75] * 4
    judge = get_metric(config.metrics[0])
    assert isinstance(judge, LLMJudge)
    assert judge.parse("**Score:** 3") == 3
    assert judge.parse("Score: 2\nOn reflection, Score: 4") == 4
    assert judge.parse("4") == 4
    assert judge.parse("Score: 9") is None and judge.parse("no idea") is None
    messages = judge.messages("task", "out {rubric}", None)
    assert "Is it polite? 1-5" in messages[-1]["content"]
    assert "out {rubric}" in messages[-1]["content"]  # placeholders in outputs stay as they are
    assert "[Reference answer]" not in messages[-1]["content"]


def test_unparsed_and_failed_judgments() -> None:
    def flaky(model: ModelSpec, prompt: str, attempt: int) -> str | Exception:
        if model.id == JUDGE:
            if "Tokyo" in prompt.split("[Response]")[1]:
                return ProviderError("judge down")
            return "I refuse to give a number."
        return respond(model, prompt, attempt)

    results = run(make_config(), DATA, backend=FakeBackend(flaky))
    good = results.scores["good"]["llm_judge"]
    assert good.value is None and good.extra["unparsed"] == 2 and good.extra["judge_failed"] == 1


def test_checks_before_running() -> None:
    no_key = {"input": "q", "id": "id"}
    with pytest.raises(ConfigError, match="needs a `rubric`"):
        run(make_config(mapping=no_key), DATA, backend=FakeBackend(respond))
    with pytest.raises(ConfigError, match="no expected-output column"):
        run(make_config({"reference": True}, mapping=no_key), DATA, backend=FakeBackend(respond))
    with pytest.raises(ConfigError, match="min_score"):
        get_metric(MetricSpec(name="llm_judge", params={"model": JUDGE, "min_score": 5}))
    with pytest.raises(ConfigError, match="provider prefix"):
        get_metric(MetricSpec(name="llm_judge", params={"model": "gpt"}))


def test_resume_reuses_judgments(tmp_path: Path) -> None:
    first = run(make_config(), DATA, backend=FakeBackend(respond))
    backend = FakeBackend(respond)
    again = run(make_config(), DATA, backend=backend, previous=first)
    assert sorted(backend.calls) == [  # only the failed rows
        ("openai/bad", "Answer: capital of Chad?"),
        ("openai/bad", "Answer: capital of Peru?"),
    ]
    assert again.scores["good"]["llm_judge"].value == 1.0

    # Changing the judge settings judges everything again.
    backend = FakeBackend(respond)
    run(make_config({"max_score": 10}), DATA, backend=backend, previous=first)
    assert sum(model_id == JUDGE for model_id, _ in backend.calls) == 5

    # A run without the judge can be judged later from its results.json.
    plain = run(make_config(metrics=["exact_match"]), DATA, backend=FakeBackend(respond))
    plain.save_json(tmp_path / "results.json")
    backend = FakeBackend(respond)
    judged = run(
        make_config(),
        DATA,
        backend=backend,
        previous=RunResults.load_json(tmp_path / "results.json"),
    )
    assert {model_id for model_id, _ in backend.calls} == {JUDGE, "openai/bad"}
    assert judged.scores["good"]["llm_judge"].value == 1.0


def test_judgments_saved_and_exported(tmp_path: Path) -> None:
    results = run(make_config(), DATA, backend=FakeBackend(respond))
    paths = write_all(results, tmp_path)
    loaded = RunResults.load_json(paths["json"])
    assert loaded.judgments["llm_judge"]["good"][0].output == "Reasoning.\nScore: 5"
    assert score(loaded)["good"]["llm_judge"].value == 1.0

    from openpyxl import load_workbook

    outputs = list(load_workbook(paths["excel"])["Outputs"].iter_rows(values_only=True))
    column = outputs[0].index("good | llm_judge judge")
    assert outputs[1][column] == "Reasoning.\nScore: 5"
    assert outputs[4][column] is None  # r4 was not judged
    assert "Judge" in paths["latex"].read_text(encoding="utf-8")


def test_scoring_without_judgments() -> None:
    results = run(make_config(metrics=[]), DATA, backend=FakeBackend(respond))
    results.config = make_config()
    assert score(results)["good"]["llm_judge"].value is None
    with pytest.raises(ConfigError, match="judge's responses"):
        get_metric(make_config().metrics[0]).compute(["x"], [["x"]])


def test_estimate_includes_judge() -> None:
    config = make_config({"model": "openai/good"})
    result = estimate(config, DATA, backend=FakeBackend(respond))
    assert len(result.judges) == 1
    judge = result.judges[0]
    assert judge.requests == 6  # 3 rows with a reference, 2 models
    assert any("favour their own answers" in w for w in judge.warnings)
    assert result.cost_usd is not None and result.cost_usd > sum(
        m.cost_usd or 0 for m in result.models
    )
    assert "judge for llm_judge" in str(result)
