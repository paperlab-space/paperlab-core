"""Load tabular data from Excel and CSV files into a :class:`Dataset`.

All cell values are converted to Unicode strings (NFC-normalised) or ``None`` for
empty cells, so downstream code never has to deal with Excel's mix of numbers,
dates, and text.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import unicodedata
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Any

from charset_normalizer import from_bytes

from paperlab_space.errors import DataError

if TYPE_CHECKING:
    from paperlab_space.data.mapping import ColumnMapping, MappedDataset

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
CSV_SUFFIXES = {".csv", ".tsv", ".txt"}

Row = dict[str, str | None]


@dataclass
class Dataset:
    """A rectangular table of string cells with named columns."""

    columns: list[str]
    rows: list[Row]
    source: str | None = None
    sheet: str | None = None
    sha256: str | None = None
    encoding: str | None = None
    warnings: list[str] = field(default_factory=list)

    def __len__(self) -> int:
        return len(self.rows)

    def column(self, name: str) -> list[str | None]:
        """Return all values of one column."""
        if name not in self.columns:
            raise DataError(
                f"Unknown column {name!r}. Available columns: {', '.join(self.columns)}"
            )
        return [row.get(name) for row in self.rows]

    def map(self, mapping: ColumnMapping) -> MappedDataset:
        """Apply a column mapping, producing records ready for prompting and scoring."""
        from paperlab_space.data.mapping import MappedDataset

        return MappedDataset.from_dataset(self, mapping)

    @classmethod
    def from_records(cls, records: Sequence[dict[str, Any]], source: str | None = None) -> Dataset:
        """Build a dataset from a list of dicts (column order follows first appearance)."""
        header: list[str] = []
        for record in records:
            for key in record:
                if key not in header:
                    header.append(key)
        columns = _clean_header(header)
        rows = [
            {col: _to_text(record.get(key)) for key, col in zip(header, columns, strict=True)}
            for record in records
        ]
        return cls(
            columns=columns,
            rows=[r for r in rows if any(v is not None for v in r.values())],
            source=source,
        )


def list_sheets(path: str | Path) -> list[str]:
    """Return the sheet names of an Excel file (a CSV file has a single unnamed sheet)."""
    path = Path(path)
    if path.suffix.lower() in CSV_SUFFIXES:
        return [path.stem]
    _check_excel(path)
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated exception types
        raise DataError(f"Could not open Excel file {path.name}: {exc}") from exc
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def load_dataset(path: str | Path, sheet: str | int | None = None) -> Dataset:
    """Load an ``.xlsx`` or ``.csv``/``.tsv`` file.

    The first non-empty row is used as the header. Fully empty rows are dropped.

    Args:
        path: File to load.
        sheet: Excel sheet name or zero-based index. Defaults to the first sheet.
            Ignored for CSV files.
    """
    path = Path(path)
    if not path.exists():
        raise DataError(f"File not found: {path}")
    suffix = path.suffix.lower()
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()

    if suffix in EXCEL_SUFFIXES:
        sheet_name, table = _read_excel(path, sheet)
        dataset = _build(table, source=str(path), sheet=sheet_name)
    elif suffix in CSV_SUFFIXES:
        text, encoding = decode_bytes(raw)
        dataset = _build(_read_csv(text, tab=suffix == ".tsv"), source=str(path))
        dataset.encoding = encoding
    elif suffix == ".xls":
        raise DataError("Legacy .xls files are not supported. Please re-save the file as .xlsx.")
    else:
        raise DataError(f"Unsupported file type {suffix!r}. Use .xlsx, .csv, or .tsv.")
    dataset.sha256 = digest
    return dataset


def decode_bytes(raw: bytes) -> tuple[str, str]:
    """Decode text bytes, trying UTF-8 first and falling back to detection."""
    for bom, encoding in (
        (b"\xef\xbb\xbf", "utf-8-sig"),
        (b"\xff\xfe", "utf-16"),
        (b"\xfe\xff", "utf-16"),
    ):
        if raw.startswith(bom):
            return raw.decode(encoding), encoding
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    best = from_bytes(raw).best()
    if best is None:
        raise DataError("Could not detect the file's text encoding. Please save it as UTF-8.")
    return str(best), best.encoding


def _check_excel(path: Path) -> None:
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        raise DataError(f"Not an Excel file: {path}")


def _read_excel(path: Path, sheet: str | int | None) -> tuple[str, list[list[Any]]]:
    _check_excel(path)
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated exception types
        raise DataError(f"Could not open Excel file {path.name}: {exc}") from exc
    try:
        names = list(workbook.sheetnames)
        if sheet is None:
            name = names[0]
        elif isinstance(sheet, int):
            if not 0 <= sheet < len(names):
                raise DataError(
                    f"Sheet index {sheet} out of range; the file has {len(names)} sheet(s)."
                )
            name = names[sheet]
        else:
            if sheet not in names:
                raise DataError(f"Sheet {sheet!r} not found. Available sheets: {', '.join(names)}")
            name = sheet
        worksheet = workbook[name]
        return name, [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_csv(text: str, tab: bool) -> list[list[Any]]:
    if tab:
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    return [list(row) for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)]


def _build(table: Iterable[Sequence[Any]], source: str, sheet: str | None = None) -> Dataset:
    rows_iter = iter(table)
    header: list[Any] | None = None
    for candidate in rows_iter:
        if any(_to_text(cell) is not None for cell in candidate):
            header = list(candidate)
            break
    if header is None:
        raise DataError("The file is empty.")

    width = len(header)
    body: list[list[str | None]] = []
    for raw_row in rows_iter:
        cells = [_to_text(cell) for cell in raw_row]
        if all(cell is None for cell in cells):
            continue
        width = max(width, len(cells))
        body.append(cells)

    # Trailing columns with neither a header nor any data are Excel noise.
    header_text = [_to_text(cell) for cell in header] + [None] * (width - len(header))
    while (
        width
        and header_text[width - 1] is None
        and all(len(row) < width or row[width - 1] is None for row in body)
    ):
        width -= 1
    columns = _clean_header(header_text[:width])

    rows: list[Row] = []
    for cells in body:
        cells = (cells + [None] * width)[:width]
        rows.append(dict(zip(columns, cells, strict=True)))
    return Dataset(columns=columns, rows=rows, source=source, sheet=sheet)


def _clean_header(header: Sequence[Any]) -> list[str]:
    """Fill blank header cells and de-duplicate names."""
    columns: list[str] = []
    seen: dict[str, int] = {}
    for index, cell in enumerate(header, start=1):
        name = _to_text(cell) or f"column_{index}"
        name = " ".join(name.split())
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
            while name in seen:
                name = f"{name}_"
        seen.setdefault(name, 1)
        columns.append(name)
    return columns


def _to_text(value: Any) -> str | None:
    """Convert a cell value to an NFC-normalised string, or ``None`` if empty."""
    if value is None:
        return None
    if isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif isinstance(value, float):
        text = str(int(value)) if value.is_integer() else repr(value)
    elif isinstance(value, dt.datetime):
        text = value.date().isoformat() if value.time() == dt.time(0) else value.isoformat()
    elif isinstance(value, (dt.date, dt.time)):
        text = value.isoformat()
    else:
        text = str(value)
    text = unicodedata.normalize("NFC", text).strip()
    return text or None
