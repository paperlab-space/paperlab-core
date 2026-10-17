"""Excel export: a summary sheet, per-row outputs and scores, and the run config."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from paperlab_core.export.repro import reproducibility
from paperlab_core.results import RunResults

EXCEL_CELL_LIMIT = 32_767
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def write_excel(results: RunResults, path: str | Path) -> Path:
    """Write ``results`` to an ``.xlsx`` workbook with Summary, Outputs, and Config sheets."""
    path = Path(path)
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Summary"
    _write_summary(summary, results)
    _write_outputs(workbook.create_sheet("Outputs"), results)
    _write_config(workbook.create_sheet("Config"), results)
    workbook.save(path)
    return path


def _write_summary(sheet: Worksheet, results: RunResults) -> None:
    metric_keys = [spec.key for spec in results.config.metrics]
    header = [
        "Model",
        "Model ID",
        *metric_keys,
        "Rows",
        "Failed",
        "Cost (USD)",
        "Input tokens",
        "Output tokens",
    ]
    _header(sheet, header)
    for name, model_run in results.runs.items():
        scores = results.scores.get(name, {})
        row: list[Any] = [name, model_run.model.id]
        row += [scores[k].value if k in scores else None for k in metric_keys]
        row += [
            len(model_run.rows),
            model_run.n_failed,
            model_run.cost_usd,
            model_run.input_tokens,
            model_run.output_tokens,
        ]
        _append(sheet, row)
    for row_cells in sheet.iter_rows(min_row=2, min_col=3, max_col=2 + len(metric_keys)):
        for cell in row_cells:
            cell.number_format = "0.0000"
    _widths(sheet, [28, 40, *[14] * len(metric_keys), 8, 8, 12, 14, 14])


def _write_outputs(sheet: Worksheet, results: RunResults) -> None:
    data = results.data
    per_row_metrics = {
        name: [k for k, r in results.scores.get(name, {}).items() if r.per_row is not None]
        for name in results.runs
    }
    has_errors = {name: any(not row.ok for row in run.rows) for name, run in results.runs.items()}

    header = ["id", *data.mapping.fields]
    if data.has_answer_key:
        header.append("expected")
    for name in results.runs:
        header.append(f"{name} | output")
        header += [f"{name} | {k}" for k in per_row_metrics[name]]
        if has_errors[name]:
            header.append(f"{name} | error")
    _header(sheet, header)

    for index, record in enumerate(data.records):
        row: list[Any] = [record.id, *[record.fields[f] for f in data.mapping.fields]]
        if data.has_answer_key:
            row.append(" | ".join(record.expected) if record.expected else None)
        for name, model_run in results.runs.items():
            output = model_run.rows[index]
            row.append(output.output)
            for key in per_row_metrics[name]:
                per_row = results.scores[name][key].per_row
                row.append(per_row[index] if per_row is not None else None)
            if has_errors[name]:
                row.append(output.error)
        _append(sheet, row)

    widths = []
    for title in header:
        is_text = (
            title in data.mapping.fields
            or title in ("expected",)
            or title.endswith(("| output", "| error"))
        )
        widths.append(45 if is_text else 12)
    _widths(sheet, widths)
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = _WRAP


def _write_config(sheet: Worksheet, results: RunResults) -> None:
    text = yaml.safe_dump(reproducibility(results), allow_unicode=True, sort_keys=False)
    for line in text.splitlines():
        _append(sheet, [line])
    sheet.column_dimensions["A"].width = 100


def _header(sheet: Worksheet, values: list[str]) -> None:
    _append(sheet, values)
    for cell in sheet[1]:
        cell.font = _BOLD
    sheet.freeze_panes = "B2"


def _append(sheet: Worksheet, values: list[Any]) -> None:
    sheet.append([_safe(v) for v in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"  # never let a model output become a formula


def _safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(value) > EXCEL_CELL_LIMIT:
        value = value[: EXCEL_CELL_LIMIT - 20] + " … [truncated]"
    return value


def _widths(sheet: Worksheet, widths: Sequence[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
